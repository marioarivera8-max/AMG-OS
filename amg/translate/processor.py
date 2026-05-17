"""Spreadsheet translation pipeline: templating, compliance, Claude polish, styled output."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from amg.translate import config
from amg.translate.claude_client import PolishRequest, polish_batch, reset_polish_limits_for_run
from amg.translate.translate_log import translate_logger
from amg.translate.schema import (
    OUTPUT_COLUMNS,
    InputRow,
    OutputRow,
    RiskLevel,
    TranslationStatus,
    validate_and_load,
)
from amg.translate.rules import (
    apply_soften_title,
    check_compliance,
    detect_archetypes,
    is_bare_produce_filename,
    merge_risk,
    merge_tags_ordered,
    pick_description,
    pick_primary_archetype,
    pick_title,
    resolve_source_language,
    studio_from_row,
    style_entry_for,
    word_count,
)


_FILL_KEEP = PatternFill("solid", fgColor="F4CCCC")
_FILL_NEW = PatternFill("solid", fgColor="D9EAD3")
_FILL_REVIEW = PatternFill("solid", fgColor="FFF2CC")
_FILL_MED = PatternFill("solid", fgColor="FFE0B2")
_FILL_HIGH = PatternFill("solid", fgColor="F8CBAD")


def _ensure_quality_notes(
    title: str,
    description: str,
    tags: str,
    *,
    check_title: bool = True,
    check_description: bool = True,
    check_tags: bool = True,
) -> list[str]:
    notes: list[str] = []
    if check_title:
        tw = word_count(title)
        if title and (tw < config.TITLE_MIN_WORDS or tw > config.TITLE_MAX_WORDS):
            notes.append(f"title_words={tw}")
    if check_description:
        dw = word_count(description)
        if description and (dw < config.DESC_MIN_WORDS or dw > config.DESC_MAX_WORDS):
            notes.append(f"description_words={dw}")
    if check_tags:
        tag_list = [t.strip() for t in tags.split(",") if t.strip()]
        if tag_list and len(tag_list) < config.TAGS_MIN:
            notes.append(f"tags_count={len(tag_list)}")
    return notes


def _row_fill(status: TranslationStatus, risk: RiskLevel) -> PatternFill:
    if risk == RiskLevel.HIGH:
        return _FILL_HIGH
    if status == TranslationStatus.NEEDS_REVIEW:
        return _FILL_REVIEW
    if risk == RiskLevel.MED:
        return _FILL_MED
    if status == TranslationStatus.NEW_TITLE:
        return _FILL_NEW
    return _FILL_KEEP


def _norm_risk(s: str) -> RiskLevel:
    u = (s or "LOW").upper()
    if u == "HIGH":
        return RiskLevel.HIGH
    if u == "MED":
        return RiskLevel.MED
    return RiskLevel.LOW


def _draft_row(inp: InputRow, *, overwrite_titles: bool) -> dict[str, Any]:
    preserve = config.preserve_existing(overwrite_titles)
    studio_cfg = studio_from_row(inp.scene_code, inp.studio or None)
    studio_label = inp.studio.strip() or studio_cfg.name
    geo = studio_cfg.geo_identifier.strip() or "Sexy"
    combined = f"{inp.source_title}\n{inp.filename}".strip()
    lang = resolve_source_language(inp.language, studio_cfg.source_language, combined)
    matches = detect_archetypes(combined, lang)
    primary = pick_primary_archetype(matches, studio_cfg.default_archetype)
    style = style_entry_for(primary.archetype)

    notes_out: list[str] = []
    review_hit = False

    keep_title = preserve and bool(inp.existing_title.strip())
    keep_desc = preserve and bool(inp.existing_description.strip())
    keep_tags = preserve and bool(inp.existing_tags.strip())

    title = inp.existing_title.strip() if keep_title else pick_title(inp.scene_code, primary.archetype, geo)
    description = (
        inp.existing_description.strip() if keep_desc else pick_description(inp.scene_code, primary.archetype, geo)
    )

    genre = list(style.get("genre_tags", ["Adult"]))
    arch_tags = list(style.get("archetype_tags", []))
    action = list(style.get("action_tags", []))
    modifier = list(style.get("modifier_tags", []))
    seen_cf: set[str] = set()
    match_tags: list[str] = []
    for t in primary.tags:
        k = t.casefold()
        if k not in seen_cf:
            seen_cf.add(k)
            match_tags.append(t)
    for m in matches[1:6]:
        for t in m.tags:
            k = t.casefold()
            if k not in seen_cf:
                seen_cf.add(k)
                match_tags.append(t)

    if keep_tags:
        tags_str = inp.existing_tags.strip()
    else:
        merged = merge_tags_ordered(
            genre=genre,
            archetype_style=arch_tags,
            match_tags=match_tags,
            action=action,
            modifier=modifier,
            geo=geo,
            limit=config.TAGS_MAX,
        )
        while len(merged) < config.TAGS_MIN:
            for pad in ("Adult", "Solo", "MILF", "Teaser"):
                if pad not in merged:
                    merged.append(pad)
                if len(merged) >= config.TAGS_MIN:
                    break
        tags_str = ", ".join(merged[: config.TAGS_MAX])

    if not keep_title:
        title, _soft_hits = apply_soften_title(title)
    else:
        _soft_hits = []

    if not keep_desc:
        description, _ = apply_soften_title(description)

    blob = " ".join([title, description, tags_str, combined]).lower()
    compliance = check_compliance(blob)
    arche_risk = primary.risk.upper() if primary.risk.upper() in {"LOW", "MED", "HIGH"} else "LOW"
    risk_s = merge_risk(arche_risk, compliance.risk_level)

    status = TranslationStatus.NEW_TITLE
    if keep_title:
        status = TranslationStatus.KEEP_ORIGINAL

    if compliance.hard_banned:
        review_hit = True
        status = TranslationStatus.NEEDS_REVIEW
        risk_s = "HIGH"
        notes_out.append("hard_ban:" + ",".join(compliance.hard_banned[:6]))

    bare = is_bare_produce_filename(inp.filename)
    if bare:
        review_hit = True
        notes_out.append("bare_filename")
        if status != TranslationStatus.NEEDS_REVIEW:
            status = TranslationStatus.NEEDS_REVIEW

    qnotes = _ensure_quality_notes(
        title,
        description,
        tags_str,
        check_title=not keep_title,
        check_description=not keep_desc,
        check_tags=not keep_tags,
    )
    if qnotes:
        notes_out.extend(qnotes)
        review_hit = True
        status = TranslationStatus.NEEDS_REVIEW

    review_col = "YES" if review_hit or risk_s == "HIGH" else ""

    return {
        "input": inp,
        "studio_name": studio_label,
        "language": lang,
        "title": title,
        "description": description,
        "tags": tags_str,
        "risk": _norm_risk(risk_s),
        "status": status,
        "review": review_col,
        "notes_out": "; ".join(notes_out),
        "primary_archetype": primary.archetype,
        "bare": bare,
        "pre_polish": (title, description, tags_str),
        "compliance_snapshot": compliance,
        "skip_claude": preserve and bool(inp.existing_title.strip()),
    }


def _eligible_for_api_polish(d: dict[str, Any]) -> bool:
    if d.get("skip_claude"):
        return False
    risk = d["risk"]
    return bool(d["bare"] or risk in (RiskLevel.MED, RiskLevel.HIGH) or d.get("review") == "YES")


def _preflight_cost_estimate(drafts: list[dict[str, Any]]) -> tuple[int, int, int, float]:
    total = len(drafts)
    will = sum(1 for d in drafts if _eligible_for_api_polish(d))
    rules = total - will
    if will <= 0:
        batches = 0
    else:
        batches = (will + config.CLAUDE_BATCH_SIZE - 1) // config.CLAUDE_BATCH_SIZE
        batches = min(batches, config.MAX_API_BATCHES_PER_RUN)
    est = round(float(batches) * float(config.ESTIMATED_COST_PER_API_BATCH_USD), 2)
    return total, rules, will, est


@dataclass(frozen=True)
class DryRunReport:
    path: Path
    row_count: int
    studio_label: str
    language_line: str
    keep_original_rows: int
    new_title_rows: int
    bare_filename_rows: int
    needs_review_rows: int
    will_polish_rows: int
    est_cost_usd: float
    est_batches: int
    med_row_indices: tuple[int, ...]
    high_row_indices: tuple[int, ...]
    outgoing_folder: str
    output_filename_pattern: str


def dry_run_analyze(input_path: Path, *, overwrite_titles: bool) -> DryRunReport:
    input_rows, _ = validate_and_load(Path(input_path))
    drafts = [_draft_row(r, overwrite_titles=overwrite_titles) for r in input_rows]
    total, _rules, will, est_cost = _preflight_cost_estimate(drafts)
    batches = 0
    if will > 0:
        batches = (will + config.CLAUDE_BATCH_SIZE - 1) // config.CLAUDE_BATCH_SIZE
        batches = min(batches, config.MAX_API_BATCHES_PER_RUN)

    keep = sum(1 for d in drafts if d["status"] == TranslationStatus.KEEP_ORIGINAL)
    new_t = sum(1 for d in drafts if d["status"] == TranslationStatus.NEW_TITLE)
    nr = sum(1 for d in drafts if d["status"] == TranslationStatus.NEEDS_REVIEW)
    bare = sum(1 for d in drafts if d["bare"])

    lang_line = str(drafts[0]["language"]) if drafts else ""
    if drafts and not drafts[0]["input"].language.strip():
        lang_line = f"{drafts[0]['language']} (from studio default)"

    studio_lbl = drafts[0]["studio_name"] if drafts else ""
    if drafts and not drafts[0]["input"].studio.strip():
        studio_lbl = f"{studio_lbl} (auto-detected)"

    stem = Path(input_path).stem
    root = config.drive_root()

    med_rows = tuple(d["input"].row_index for d in drafts if d["risk"] == RiskLevel.MED)
    high_rows = tuple(d["input"].row_index for d in drafts if d["risk"] == RiskLevel.HIGH)

    return DryRunReport(
        path=Path(input_path),
        row_count=total,
        studio_label=studio_lbl or "(unknown)",
        language_line=str(lang_line),
        keep_original_rows=keep,
        new_title_rows=new_t,
        bare_filename_rows=bare,
        needs_review_rows=nr,
        will_polish_rows=will,
        est_cost_usd=float(est_cost),
        est_batches=batches,
        med_row_indices=med_rows,
        high_row_indices=high_rows,
        outgoing_folder=str(root / config.OUTPUT_FOLDER),
        output_filename_pattern=f"{stem}_TRANSLATED_<timestamp>.xlsx",
    )


def _apply_polished(
    draft: dict[str, Any],
    scene_code: str,
    title: str,
    description: str,
    tags_list: list[str],
) -> dict[str, Any]:
    tags_str = ", ".join(tags_list[: config.TAGS_MAX])
    title, _ = apply_soften_title(title)
    description, _ = apply_soften_title(description)
    inp: InputRow = draft["input"]
    combined = f"{inp.source_title}\n{inp.filename}".strip().lower()
    blob = " ".join([title, description, tags_str, combined])
    compliance = check_compliance(blob)
    risk_s = merge_risk(draft["risk"].value, compliance.risk_level)
    if compliance.hard_banned:
        # revert
        t0, d0, tag0 = draft["pre_polish"]
        draft["title"] = t0
        draft["description"] = d0
        draft["tags"] = tag0
        draft["risk"] = RiskLevel.HIGH
        draft["status"] = TranslationStatus.NEEDS_REVIEW
        draft["review"] = "YES"
        extra = draft["notes_out"]
        draft["notes_out"] = (extra + "; polish_reverted_hard_ban").strip("; ")
        return draft
    draft["title"] = title
    draft["description"] = description
    draft["tags"] = tags_str
    draft["risk"] = _norm_risk(risk_s)
    if draft["bare"]:
        draft["status"] = TranslationStatus.NEEDS_REVIEW
        draft["review"] = "YES"
    qnotes = _ensure_quality_notes(title, description, tags_str)
    if qnotes:
        draft["status"] = TranslationStatus.NEEDS_REVIEW
        draft["review"] = "YES"
        draft["notes_out"] = (draft["notes_out"] + "; " + ";".join(qnotes)).strip("; ")
    return draft


def _run_claude_pass(drafts: list[dict[str, Any]], *, enabled: bool) -> None:
    if not enabled:
        return
    batch: list[tuple[int, PolishRequest]] = []
    for idx, d in enumerate(drafts):
        if d.get("skip_claude"):
            continue
        risk = d["risk"]
        if d["bare"] or risk in (RiskLevel.MED, RiskLevel.HIGH) or d.get("review") == "YES":
            inp = d["input"]
            tags_list = [t.strip() for t in d["tags"].split(",") if t.strip()]
            batch.append(
                (
                    idx,
                    PolishRequest(
                        scene_code=inp.scene_code,
                        source_title=inp.source_title,
                        filename=inp.filename,
                        studio=d["studio_name"],
                        archetype=d["primary_archetype"],
                        draft_title=d["title"],
                        draft_description=d["description"],
                        tags=tags_list,
                        risk=risk.value,
                    ),
                )
            )
    for i in range(0, len(batch), config.CLAUDE_BATCH_SIZE):
        chunk = batch[i : i + config.CLAUDE_BATCH_SIZE]
        results = polish_batch([c[1] for c in chunk])
        by_code = {r.scene_code: r for r in results}
        for idx, req in chunk:
            hit = by_code.get(req.scene_code)
            if not hit:
                continue
            drafts[idx] = _apply_polished(
                drafts[idx],
                req.scene_code,
                hit.title,
                hit.description,
                hit.tags,
            )


def process_workbook(
    input_path: Path,
    output_path: Path | None,
    *,
    overwrite_titles: bool,
    use_claude: bool,
) -> Path:
    reset_polish_limits_for_run()
    input_rows, _sheet_title = validate_and_load(Path(input_path))
    drafts = [_draft_row(r, overwrite_titles=overwrite_titles) for r in input_rows]

    rows_n, rules_only, polish_n, est_cost = _preflight_cost_estimate(drafts)
    translate_logger().info(
        "FILE: %s | rows=%d | rules_only=%d | will_polish=%d | est_cost=$%.2f",
        Path(input_path).name,
        rows_n,
        rules_only,
        polish_n,
        est_cost,
    )

    claude_on = bool(use_claude and config.anthropic_enabled())
    _run_claude_pass(drafts, enabled=claude_on)

    outputs: list[OutputRow] = []
    for d in drafts:
        inp = d["input"]
        outputs.append(
            OutputRow(
                scene_code=inp.scene_code,
                source_title=inp.source_title,
                filename=inp.filename,
                studio=d["studio_name"],
                language=d["language"],
                title=d["title"],
                description=d["description"],
                tags=d["tags"],
                risk=d["risk"],
                review=d["review"],
                status=d["status"],
                notes_in=inp.notes,
                notes_out=d["notes_out"],
            )
        )

    dest = output_path or Path(input_path).with_name(f"{Path(input_path).stem}_TRANSLATED.xlsx")
    write_output_workbook(outputs, dest)
    return dest


def write_output_workbook(rows: list[OutputRow], dest: Path) -> None:
    wb = Workbook()
    banner = wb.active
    banner.title = "!! READ ME FIRST !!"
    banner["A1"] = "TRANSLATE MODULE — READ ME FIRST"
    banner["A1"].font = Font(size=14, bold=True)
    lines = [
        "Row colors (Translated sheet):",
        f"  Pink/red fill F4CCCC = {TranslationStatus.KEEP_ORIGINAL.value}",
        f"  Green fill D9EAD3 = {TranslationStatus.NEW_TITLE.value}",
        f"  Yellow fill FFF2CC = {TranslationStatus.NEEDS_REVIEW.value}",
        "  Orange fill FFE0B2 = MED compliance risk",
        "  Peach fill F8CBAD = HIGH compliance risk",
        "",
        f"Total rows exported: {len(rows)}",
        "Columns: scene metadata plus English title, description, comma-separated tags, risk, review flag, status, notes.",
    ]
    for i, line in enumerate(lines, start=3):
        banner.cell(row=i, column=1, value=line)

    ws = wb.create_sheet("Translated")
    _hdr_fill = PatternFill("solid", fgColor="4A4A4A")
    _hdr_font = Font(bold=True, color="FFFFFF")
    for col, head in enumerate(OUTPUT_COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=head)
        c.font = _hdr_font
        c.fill = _hdr_fill
    for r_idx, row in enumerate(rows, start=2):
        values = [
            row.scene_code,
            row.source_title,
            row.filename,
            row.studio,
            row.language,
            row.title,
            row.description,
            row.tags,
            row.risk.value,
            row.review,
            row.status.value,
            row.notes_in,
            row.notes_out,
        ]
        fill = _row_fill(row.status, row.risk)
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for col in range(1, len(OUTPUT_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["G"].width = 36
    ws.column_dimensions["H"].width = 36

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Metric"
    summary["B1"] = "Count"
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    ctr_status = Counter(r.status.value for r in rows)
    ctr_risk = Counter(r.risk.value for r in rows)
    r_i = 2
    summary.cell(row=r_i, column=1, value="Rows")
    summary.cell(row=r_i, column=2, value=len(rows))
    r_i += 1
    for label, ctr in (("Status", ctr_status), ("Risk", ctr_risk)):
        summary.cell(row=r_i, column=1, value=label)
        r_i += 1
        for k, v in sorted(ctr.items()):
            summary.cell(row=r_i, column=1, value=k)
            summary.cell(row=r_i, column=2, value=v)
            r_i += 1

    wb.save(dest)


def write_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Catalog"
    headers = [
        "scene_code",
        "source_title",
        "filename",
        "studio",
        "language",
        "notes",
        "title",
        "description",
        "tags",
    ]
    for idx, head in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=head)
        ws.cell(row=1, column=idx).font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(path)
