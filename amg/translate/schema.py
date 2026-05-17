"""Input/output sheet schema."""
from __future__ import annotations

import re
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

import openpyxl

from amg.translate import config


def _normalize_header_value(raw: object) -> str:
    """Lowercase; collapse spaces/hyphens to underscores (Scene Code → scene_code)."""
    if raw is None:
        return ""
    s = str(raw).strip().lower()
    if not s:
        return ""
    s = re.sub(r"[\s\-]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s


_HEADER_SYNONYMS = {
    "scenecode": "scene_code",
    "scene_id": "scene_code",
    "id_scene": "scene_code",
    "source_title_original": "source_title",
    "original_title": "source_title",
    "src_title": "source_title",
    "file_name": "filename",
    "media_filename": "filename",
}


def _canonical_header(norm: str) -> str:
    return _HEADER_SYNONYMS.get(norm, norm)

REQUIRED_COLUMNS = ["scene_code"]
KNOWN_COLUMNS = [
    "scene_code",
    "source_title",
    "filename",
    "studio",
    "language",
    "notes",
    "title",
    "description",
    "tags",
]

OUTPUT_COLUMNS = [
    "scene_code",
    "source_title",
    "filename",
    "studio",
    "language",
    "title",
    "description",
    "tags",
    "risk",
    "review",
    "status",
    "notes_in",
    "notes_out",
]


class SchemaError(Exception):
    pass


class RiskLevel(str, Enum):
    LOW = "LOW"
    MED = "MED"
    HIGH = "HIGH"


class TranslationStatus(str, Enum):
    KEEP_ORIGINAL = "KEEP ORIGINAL"
    NEW_TITLE = "NEW TITLE"
    NEEDS_REVIEW = "NEEDS REVIEW"


@dataclass
class InputRow:
    row_index: int
    scene_code: str
    source_title: str = ""
    filename: str = ""
    studio: str = ""
    language: str = ""
    notes: str = ""
    existing_title: str = ""
    existing_description: str = ""
    existing_tags: str = ""
    extra: dict[str, str] | None = None


@dataclass
class OutputRow:
    scene_code: str
    source_title: str
    filename: str
    studio: str
    language: str
    title: str
    description: str
    tags: str
    risk: RiskLevel
    review: str
    status: TranslationStatus
    notes_in: str
    notes_out: str


def validate_and_load(xlsx_path: Path) -> tuple[list[InputRow], str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    if "Catalog" in wb.sheetnames:
        ws = wb["Catalog"]
    headers = [_canonical_header(_normalize_header_value(c.value)) for c in ws[1]]
    if not any(headers):
        raise SchemaError(f"{xlsx_path.name}: header row (row 1) is empty")
    col_idx = {h: i for i, h in enumerate(headers) if h}
    missing = [c for c in REQUIRED_COLUMNS if c not in col_idx]
    if missing:
        raise SchemaError(f"{xlsx_path.name}: missing required column(s): {missing}")

    rows: list[InputRow] = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue

        def get(col_name: str) -> str:
            i = col_idx.get(col_name)
            if i is None or i >= len(row):
                return ""
            v = row[i]
            return "" if v is None else str(v).strip()

        scene_code = get("scene_code")
        if not scene_code:
            continue
        extra: dict[str, str] = {}
        for h, i in col_idx.items():
            if h not in KNOWN_COLUMNS and i < len(row):
                v = row[i]
                if v is not None and str(v).strip():
                    extra[h] = str(v).strip()
        rows.append(
            InputRow(
                row_index=row_num,
                scene_code=scene_code,
                source_title=get("source_title"),
                filename=get("filename"),
                studio=get("studio"),
                language=get("language").lower(),
                notes=get("notes"),
                existing_title=get("title"),
                existing_description=get("description"),
                existing_tags=get("tags"),
                extra=extra or None,
            )
        )
    if not rows:
        raise SchemaError(f"{xlsx_path.name}: no data rows found")
    if len(rows) > config.MAX_ROWS_PER_FILE:
        raise SchemaError(
            f"{xlsx_path.name}: row count {len(rows)} exceeds limit {config.MAX_ROWS_PER_FILE} "
            f"(configure AMG_TRANSLATE_MAX_ROWS)\nBatch must be split into smaller spreadsheets."
        )
    return rows, ws.title
