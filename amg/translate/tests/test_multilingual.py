"""Multilingual archetype coverage + universal hotwords."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from amg.translate.processor import process_workbook
from amg.translate.rules import detect_language


def _hdr_map(ws):
    return {
        str(c.value).strip(): idx
        for idx, c in enumerate(ws[1], start=1)
        if c.value is not None and str(c.value).strip()
    }


_CASES: list[tuple[str, str, str, str, str]] = [
    # scene_code, lang_cell, source_title, filename, substring_in_title
    ("MG_ES_SM", "es", "Madrastra traviesa en casa", "mg_es_sm.mp4", "Stepmom"),
    ("MG_ES_SEC", "es", "Secretaria traviesa despues de hora", "mg_es_sec.mp4", "Secretary"),
    ("RB_IT_SM", "it", "Matrigna italiana gioca con te", "rb_it_sm.mp4", "Stepmom"),
    ("RB_IT_SEC", "it", "Segretaria vuole il bonus", "rb_it_sec.mp4", "Secretary"),
    ("MG_PT_SM", "pt", "Madrasta sedenta por atencao", "mg_pt_sm.mp4", "Stepmom"),
    ("MG_PT_SEC", "pt", "Secretaria manda no escritorio", "mg_pt_sec.mp4", "Secretary"),
    ("YBR_FR_SM", "fr", "Belle-mère surveille ta douche", "ybr_fr_sm.mp4", "Stepmom"),
    ("YBR_FR_SEC", "fr", "Secrétaire bruyante au bureau", "ybr_fr_sec.mp4", "Secretary"),
    ("BH_DE_SM", "de", "stiefmutter fantasy nach mitternacht", "bh_de_sm.mp4", "Stepmom"),
    ("BH_DE_SEC", "de", "sekretärin bleibt spontan da", "bh_de_sec.mp4", "Secretary"),
    ("SHN_JA_SM", "ja", "義母との昼下がりドキュメント", "shn_ja_sm.mp4", "Stepmom"),
    ("SHN_JA_SEC", "ja", "秘書が残業で大胆になる", "shn_ja_sec.mp4", "Secretary"),
    ("SHN_JA_ROM", "ja", "gibo summer heat diary clip", "shn_ja_rom.mp4", "Stepmom"),
    ("YB_KO_SM", "ko", "새엄마와 단둘이 남은 오후", "yb_ko_sm.mp4", "Stepmom"),
    ("YB_KO_SEC", "ko", "비서가 회의 후 단단히 부탁한다", "yb_ko_sec.mp4", "Secretary"),
    ("YB_ZH_SM", "zh", "继母在厨房里低声挑逗", "yb_zh_sm.mp4", "Stepmom"),
    ("YB_ZH_SEC", "zh", "秘书加班后大胆撒娇", "yb_zh_sec.mp4", "Secretary"),
    ("MG_JOI", "es", "cuenta regresiva pov mirando camara", "joi_countdown.mp4", "MILF"),
    ("BH_CP_EXTRA", "de", "bonus splash clip", "creampie_bonus_bh.mp4", "Creampie"),
    ("BH_ANAL", "de", "stretch session privat", "anal_training_bh.mp4", "Anal"),
    ("BH_SQ", "de", "splash tutorial für dich", "squirt_show_bh.mp4", "Squirt"),
    ("BH_MIX_JOI", "de", "nothing spicy here", "solo_joi_tease.mp4", "MILF"),
    ("ZZ_DETECT_ES", "", "madrastra necesita ayuda urgente", "zz_detect.mp4", "Stepmom"),
    ("ZZ_DETECT_DE", "", "stiefschwester test clip neutral", "zz_de.mp4", "Stepsister"),
    ("MG_CE_SOLO", "es", "latina solo mirando lente", "mg_ce_solo.mp4", "Latino"),
    ("RB_FEM_ES", "es", "femdom countdown latino office", "rb_fem.mp4", "Femdom"),
    ("BH_CUCK", "de", "hotwife erzählt beim date", "cuckold_story_bh.mp4", "Cuckold"),
    ("BH_LTX", "de", "geschmeidiges outfit clip", "latex_shine_bh.mp4", "Latex"),
    ("BH_FT", "de", "zehenlack pov strip", "foot_focus_bh.mp4", "Foot"),
    ("BH_EDG", "de", "slow stroke coaching mp4", "edging_coach_bh.mp4", "Edging"),
    ("BH_SELF", "de", "mirror smartphone clip", "selfie_mirror_bh.mp4", "Selfie"),
]


class MultilingualTests(unittest.TestCase):
    def setUp(self) -> None:
        os.environ["AMG_TRANSLATE_CLAUDE"] = "0"

    def test_thirty_plus_language_pack_coverage(self) -> None:
        self.assertGreaterEqual(len(_CASES), 30)

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp_path = Path(tmp.name)
        tmp.close()

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Catalog"
        headers = [
            "studio",
            "scene_code",
            "language",
            "source_title",
            "filename",
            "title",
            "description",
            "tags",
        ]
        for col, head in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=head)

        for scene_code, lang, src_title, fn, _needle in _CASES:
            ws.append(["", scene_code, lang, src_title, fn, "", "", ""])

        wb.save(tmp_path)
        output = tmp_path.with_name(tmp_path.stem + "_out.xlsx")
        process_workbook(tmp_path, output, overwrite_titles=False, use_claude=False)

        ws_out = load_workbook(output)["Translated"]
        h = _hdr_map(ws_out)

        by_code = {
            str(ws_out.cell(row=r, column=h["scene_code"]).value): r
            for r in range(2, ws_out.max_row + 1)
        }

        for scene_code, lang_cell, src_title, fn, needle in _CASES:
            row_idx = by_code[scene_code]
            resolved_lang = str(ws_out.cell(row=row_idx, column=h["language"]).value or "")
            title = str(ws_out.cell(row=row_idx, column=h["title"]).value or "")
            tags = str(ws_out.cell(row=row_idx, column=h["tags"]).value or "")
            self.assertIn(
                needle,
                title,
                msg=f"{scene_code}: expected {needle!r} in title {title!r}",
            )
            self.assertTrue(tags.strip(), msg=f"{scene_code}: tags missing")

            if lang_cell:
                self.assertEqual(resolved_lang, lang_cell)

        for code in ("MG_JOI", "BH_MIX_JOI"):
            tags_j = str(ws_out.cell(row=by_code[code], column=h["tags"]).value or "")
            self.assertIn("JOI", tags_j.upper(), msg=f"{code} missing JOI tag")

        tmp_path.unlink(missing_ok=True)
        output.unlink(missing_ok=True)

    def test_auto_detect_prefers_clear_language_signal(self) -> None:
        self.assertEqual(detect_language("madrastra culona en casa"), "es")
        self.assertEqual(detect_language("stiefbruder und stiefschwester teaser"), "de")


if __name__ == "__main__":
    unittest.main()
