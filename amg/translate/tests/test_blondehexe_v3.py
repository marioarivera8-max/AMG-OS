"""Regression suite for BlondeHexe (German) routing and catalog-quality gates."""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from amg.translate.processor import process_workbook
from amg.translate.rules import word_count


def _hdr_map(ws):
    return {
        str(c.value).strip(): idx
        for idx, c in enumerate(ws[1], start=1)
        if c.value is not None and str(c.value).strip()
    }


def _cell(ws, row: int, name: str, h: dict[str, int]):
    return ws.cell(row=row, column=h[name]).value


# 25+ realistic-ish German inventory rows (codes illustrative).
_BLONDEHexe_ROWS: list[tuple[str, str, str]] = [
    ("BH_VR099", "Besser als Deine Ehefrau", "bh_vr099_uhd.mp4"),
    ("BH_Solo335", "solo workout teaser", "zitterorg.mp4"),
    ("BH_Solo302", "", "stiefmuklein.mp4"),
    ("BH_Solo224", "produce batch", "Produce_67.mp4"),
    ("BH_GH_001", "Stiefmutter erwischt dich beim wichsen", "bh_gh_001.mp4"),
    ("BH_OFC_014", "sekretärin bleibt nach feier da", "bh_ofc_014.mp4"),
    ("BH_KNK_050", "latex catsuit dehnübung", "bh_knk_050.mp4"),
    ("BH_FEET_011", "footjob unter dem schreibtisch", "bh_feet_011.mp4"),
    ("BH_ANAL_009", "anal training für milfs", "bh_anal_009.mp4"),
    ("BH_YOGA_033", "yoga pants durchgeschwitzt", "bh_yoga_033.mp4"),
    ("BH_CAM_044", "livecam countdown session", "bh_cam_044.mp4"),
    ("BH_JOI_120", "joi countdown german dirty talk", "bh_joi_120.mp4"),
    ("BH_SQ_077", "squirt explosion auf leather couch", "bh_sq_077.mp4"),
    ("BH_CP_088", "creampie surprise nach date night", "bh_cp_088.mp4"),
    ("BH_WET_019", "wetlook dusche mit neonlicht", "bh_wet_019.mp4"),
    ("BH_EDG_031", "edging marathon für dich", "bh_edg_031.mp4"),
    ("BH_DIR_212", "dirndl strip nach oktoberfest", "bh_dir_212.mp4"),
    ("BH_LEH_055", "lehrerin gibt nachhilfe zuhause", "bh_leh_055.mp4"),
    ("BH_FLT_066", "stewardess fantasy mile high tease", "bh_flt_066.mp4"),
    ("BH_POL_072", "polizistin uniform fantasy", "bh_pol_072.mp4"),
    ("BH_MAID_083", "dienstmädchen putzt mehr als staub", "bh_maid_083.mp4"),
    ("BH_NEI_091", "nachbarin flirtet im treppenhaus", "bh_nei_091.mp4"),
    ("BH_GYM_099", "sportlehrerin schwitzt im fitnessraum", "bh_gym_099.mp4"),
    ("BH_THER_103", "therapeutin hilft beim stressabbau", "bh_ther_103.mp4"),
    ("BH_BOSS_111", "chefin fordert extras nach meeting", "bh_boss_111.mp4"),
    ("BH_PARTY_125", "party girl nach club mit selfies", "bh_party_125.mp4"),
    ("BH_BAR_130", "bar flirt mit großen tits teaser", "bh_bar_130.mp4"),
]


class BlondeHexeTests(unittest.TestCase):
    def setUp(self) -> None:
        os.environ["AMG_TRANSLATE_CLAUDE"] = "0"

    def test_catalog_quality_and_flags(self) -> None:
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp_path = Path(tmp.name)
        tmp.close()

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Catalog"
        headers = [
            "studio",
            "scene_code",
            "language",
            "source_title",
            "filename",
            "title",
            "description",
            "tags",
        ]
        for col, head in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=head)

        for scene_code, src_title, fn in _BLONDEHexe_ROWS:
            ws.append(["", scene_code, "", src_title, fn, "", "", ""])

        wb.save(tmp_path)
        output = tmp_path.with_name(tmp_path.stem + "_out.xlsx")
        process_workbook(tmp_path, output, overwrite_titles=False, use_claude=False)

        self.assertGreaterEqual(len(_BLONDEHexe_ROWS), 25)

        out_wb = load_workbook(output)
        ws_out = out_wb["Translated"]
        h = _hdr_map(ws_out)

        produce_review = False
        stepfamily_med = False

        for row_idx in range(2, ws_out.max_row + 1):
            code = _cell(ws_out, row_idx, "scene_code", h)
            title = str(_cell(ws_out, row_idx, "title", h) or "")
            desc = str(_cell(ws_out, row_idx, "description", h) or "")
            tags = str(_cell(ws_out, row_idx, "tags", h) or "")
            status = str(_cell(ws_out, row_idx, "status", h) or "")
            risk = str(_cell(ws_out, row_idx, "risk", h) or "")
            review = str(_cell(ws_out, row_idx, "review", h) or "")
            notes = str(_cell(ws_out, row_idx, "notes_out", h) or "").lower()

            self.assertTrue(desc.strip(), f"{code}: missing description")
            self.assertTrue(tags.strip(), f"{code}: missing tags")

            if status != "KEEP ORIGINAL":
                tw = word_count(title)
                self.assertGreaterEqual(
                    tw,
                    5,
                    f"{code}: title too short ({tw} words): {title!r}",
                )
                self.assertLessEqual(
                    tw,
                    10,
                    f"{code}: title too long ({tw} words): {title!r}",
                )

            if code == "BH_Solo224":
                self.assertEqual(status, "NEEDS REVIEW")
                self.assertIn("bare_filename", notes.replace(" ", ""))
                produce_review = True

            if "stepmom" in tags.lower() or "stepmom" in title.lower():
                self.assertEqual(risk, "MED")
                stepfamily_med = True

            if "Stiefmutter" in str(
                _cell(ws_out, row_idx, "source_title", h) or ""
            ) or "stiefmu" in str(_cell(ws_out, row_idx, "filename", h) or "").lower():
                self.assertEqual(risk, "MED")

        self.assertTrue(produce_review)
        self.assertTrue(stepfamily_med)

        tmp_path.unlink(missing_ok=True)
        output.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
