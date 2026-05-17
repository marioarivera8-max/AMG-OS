"""Preservation + overwrite behaviour."""

from __future__ import annotations

import importlib
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import amg.translate.config as translate_config
from amg.translate.processor import process_workbook


def _hdr_map(ws):
    return {
        str(c.value).strip(): idx
        for idx, c in enumerate(ws[1], start=1)
        if c.value is not None and str(c.value).strip()
    }


class PreservationTests(unittest.TestCase):
    def setUp(self) -> None:
        os.environ["AMG_TRANSLATE_CLAUDE"] = "0"

    def test_preservation_keeps_populated_cells(self) -> None:
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp_path = Path(tmp.name)
        tmp.close()

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Catalog"
        headers = [
            "studio",
            "scene_code",
            "language",
            "source_title",
            "filename",
            "title",
            "description",
            "tags",
        ]
        for col, head in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=head)

        ws.append(
            [
                "BlondeHexe",
                "BH_KEEP_1",
                "de",
                "irrelevant german filler text",
                "file.mp4",
                "Existing English Title",
                "Existing English description stays.",
                "MILF, Solo",
            ]
        )
        ws.append(
            [
                "",
                "BH_NEW_2",
                "",
                "sekretärin fantasy nach feier",
                "bh_new_2.mp4",
                "",
                "",
                "",
            ]
        )
        ws.append(
            [
                "",
                "BH_BARE_3",
                "",
                "",
                "Produce_12.mp4",
                "",
                "",
                "",
            ]
        )
        wb.save(tmp_path)

        output = tmp_path.with_name(tmp_path.stem + "_out.xlsx")
        process_workbook(tmp_path, output, overwrite_titles=False, use_claude=False)

        ws_out = load_workbook(output)["Translated"]
        h = _hdr_map(ws_out)

        self.assertEqual(ws_out.cell(row=2, column=h["title"]).value, "Existing English Title")
        self.assertEqual(
            ws_out.cell(row=2, column=h["description"]).value,
            "Existing English description stays.",
        )
        self.assertEqual(ws_out.cell(row=2, column=h["tags"]).value, "MILF, Solo")
        self.assertEqual(ws_out.cell(row=2, column=h["status"]).value, "KEEP ORIGINAL")

        self.assertEqual(ws_out.cell(row=3, column=h["status"]).value, "NEW TITLE")

        self.assertEqual(ws_out.cell(row=4, column=h["status"]).value, "NEEDS REVIEW")

        tmp_path.unlink(missing_ok=True)
        output.unlink(missing_ok=True)

    def test_overwrite_env_disables_preservation(self) -> None:
        os.environ["AMG_TRANSLATE_OVERWRITE"] = "1"
        importlib.reload(translate_config)

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp_path = Path(tmp.name)
        tmp.close()

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Catalog"
        headers = [
            "studio",
            "scene_code",
            "language",
            "source_title",
            "filename",
            "title",
            "description",
            "tags",
        ]
        for col, head in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=head)

        ws.append(
            [
                "BlondeHexe",
                "BH_OW_1",
                "de",
                "latex fantasy studio session",
                "bh_ow_1.mp4",
                "Old Cached Title",
                "Old Cached Body Copy Here.",
                "Old, Tags",
            ]
        )
        wb.save(tmp_path)

        output = tmp_path.with_name(tmp_path.stem + "_overwrite.xlsx")
        process_workbook(tmp_path, output, overwrite_titles=False, use_claude=False)

        ws_out = load_workbook(output)["Translated"]
        h = _hdr_map(ws_out)
        status = ws_out.cell(row=2, column=h["status"]).value
        title = str(ws_out.cell(row=2, column=h["title"]).value or "")
        self.assertNotEqual(status, "KEEP ORIGINAL")
        self.assertNotEqual(title, "Old Cached Title")

        tmp_path.unlink(missing_ok=True)
        output.unlink(missing_ok=True)

        os.environ.pop("AMG_TRANSLATE_OVERWRITE", None)
        importlib.reload(translate_config)


if __name__ == "__main__":
    unittest.main()
