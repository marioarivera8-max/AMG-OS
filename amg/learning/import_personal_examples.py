"""
Import operator-labeled personal examples into AMG training store.

Accepted formats:
- CSV with header row
- JSONL (one object per line)
- XLSX workbooks (including embedded hyperlinks)

Canonical output written to:
  data/training/examples/<dataset_name>.jsonl
"""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional

from amg.config import TRAINING_EXAMPLES_DIR


VALID_DECISIONS = {"keep", "maybe", "reject"}


@dataclass
class ImportStats:
    input_rows: int = 0
    accepted_rows: int = 0
    skipped_rows: int = 0
    errors: int = 0
    deduped_rows: int = 0
    output_path: Optional[Path] = None


def import_personal_examples(
    input_path: Path,
    dataset_name: str = "personal_examples",
    dry_run: bool = False,
) -> ImportStats:
    stats = ImportStats()
    rows = list(_read_rows(input_path))
    stats.input_rows = len(rows)

    normalized: List[dict] = []
    seen_keys: set[str] = set()
    for idx, row in enumerate(rows, start=1):
        try:
            out = _normalize_row(row, idx)
            if out is None:
                stats.skipped_rows += 1
                continue
            dedupe_key = _dedupe_key(out)
            if dedupe_key in seen_keys:
                stats.deduped_rows += 1
                continue
            seen_keys.add(dedupe_key)
            normalized.append(out)
            stats.accepted_rows += 1
        except Exception:
            stats.errors += 1

    if dry_run:
        return stats

    TRAINING_EXAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    out_path = TRAINING_EXAMPLES_DIR / f"{_safe_name(dataset_name)}.jsonl"
    with open(out_path, "w") as f:
        for row in normalized:
            f.write(json.dumps(row) + "\n")
    stats.output_path = out_path
    return stats


def _read_rows(path: Path) -> Iterable[Dict[str, Any]]:
    p = Path(path)
    if not p.exists():
        return []
    ext = p.suffix.lower()
    if ext == ".xlsx":
        return _read_xlsx(p)
    if p.suffix.lower() == ".jsonl":
        return _read_jsonl(p)
    return _read_csv(p)


def _read_csv(path: Path) -> Iterable[Dict[str, Any]]:
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            yield dict(row)


def _read_jsonl(path: Path) -> Iterable[Dict[str, Any]]:
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            obj = json.loads(line)
            if isinstance(obj, dict):
                yield obj


def _read_xlsx(path: Path) -> Iterable[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("openpyxl is required to import .xlsx files") from e

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    for ws in wb.worksheets:
        header_row_idx = _detect_header_row(ws)
        headers = _row_headers(ws, header_row_idx)
        if not headers:
            continue
        for row_idx in range(header_row_idx + 1, ws.max_row + 1):
            cells = [ws.cell(row=row_idx, column=c) for c in range(1, ws.max_column + 1)]
            values = [c.value for c in cells]
            if _is_empty_row(values):
                continue

            row: Dict[str, Any] = {
                "_source_sheet": ws.title,
                "_source_row": row_idx,
                "_source_file": path.name,
            }
            for i, header in enumerate(headers):
                if i >= len(cells):
                    break
                if not header:
                    continue
                cell = cells[i]
                value = cell.value
                if value is None:
                    continue
                row[header] = str(value).strip()
                if cell.hyperlink is not None and cell.hyperlink.target:
                    row[f"{header}__link"] = str(cell.hyperlink.target).strip()
            yield row


def _normalize_row(row: Dict[str, Any], row_num: int) -> Optional[dict]:
    scene_id = _first(row, "scene_id", "scene", "source_scene", "scene_code", "scene_number")
    filename = _first(row, "filename", "image", "image_file", "file", "file_name", "content")
    image_path = _first(row, "image_path", "path", "file_path", "screengrabs", "cover_art")
    scene_title = _first(row, "title", "scene_name", "scene_title", "scene")
    description = _first(row, "description", "synopsis")
    if not scene_id and not filename and not image_path and not scene_title and not description:
        return None

    operator_score = _normalize_score(_first(row, "operator_score", "score", "your_score"))
    model_score = _normalize_score(_first(row, "model_score", "ai_score"))
    decision = (_first(row, "decision", "label", "operator_decision") or "").strip().lower()
    if decision and decision not in VALID_DECISIONS:
        decision = ""
    if not decision:
        decision = _derive_decision_from_status(row)

    performers = _extract_performers(row)
    categories = _normalize_list(_first(row, "categories", "category", "genre"))
    tags = _normalize_list(_first(row, "tags", "keywords", "additional_tags"))
    content_links = _collect_links(row)
    notes = (_first(row, "notes", "reason", "remark", "info") or "").strip() or None
    studio = (_first(row, "studio", "company", "creator", "creators___studios", "_source_sheet") or "").strip() or None

    normalized = {
        "source_type": "personal_example",
        "source_channel": "xlsx_tracker" if str(row.get("_source_file", "")).lower().endswith(".xlsx") else "manual",
        "source_file": row.get("_source_file"),
        "source_sheet": row.get("_source_sheet"),
        "source_row_num": row.get("_source_row"),
        "row_num": row_num,
        "scene_id": (scene_id or "").strip() or _fallback_scene_id(row),
        "filename": (filename or "").strip() or None,
        "image_path": (image_path or "").strip() or None,
        "title": (scene_title or "").strip() or None,
        "description": (description or "").strip() or None,
        "studio": studio,
        "performers": performers,
        "categories": categories,
        "tags": tags,
        "notes": notes,
        "model_score_100": model_score,
        "operator_score_100": operator_score,
        "decision": decision or None,
        "should_have_picked": (_first(row, "should_have_picked", "replacement", "better_frame") or "").strip() or None,
        "source_links": content_links,
        "record_quality": _record_quality(
            scene_title=scene_title,
            description=description,
            performers=performers,
            categories=categories,
            tags=tags,
            links=content_links,
            decision=decision,
            score=operator_score,
        ),
    }

    if _is_low_signal(normalized):
        return None
    return normalized


def _first(row: Dict[str, Any], *keys: str) -> Optional[str]:
    normalized = {_canonical_key(k): v for k, v in row.items()}
    for k in keys:
        ck = _canonical_key(k)
        if ck in normalized and normalized[ck] not in (None, ""):
            return str(normalized[ck])
    return None


def _normalize_list(raw: Optional[str]) -> List[str]:
    if not raw:
        return []
    tokens = [t.strip() for t in str(raw).split(",")]
    out = [t for t in tokens if t]
    return list(dict.fromkeys(out))


def _extract_performers(row: Dict[str, Any]) -> List[str]:
    values: List[str] = []
    for k, v in row.items():
        ck = _canonical_key(k)
        if ck.startswith("talent") or ck in {"performer", "performers", "main_talent", "cast"}:
            if v not in (None, ""):
                values.append(str(v))
    if not values:
        single = _first(row, "performers", "performer", "cast")
        if single:
            values.append(single)
    merged = ",".join(values)
    return _normalize_list(merged)


def _collect_links(row: Dict[str, Any]) -> List[str]:
    links: List[str] = []
    for k, v in row.items():
        if k.endswith("__link") and v:
            links.append(str(v).strip())
        elif isinstance(v, str) and v.startswith(("http://", "https://", "ftp://")):
            links.append(v.strip())
    # Keep order, remove duplicates.
    return list(dict.fromkeys([x for x in links if x]))


def _derive_decision_from_status(row: Dict[str, Any]) -> str:
    status_markers = 0
    for k, v in row.items():
        ck = _canonical_key(k)
        if ck in {"adult_empire", "aebn", "ave_ent", "sexlikereal", "veegaz", "faphouse"}:
            text = str(v or "").strip().lower()
            if text and text not in {"", "no", "n", "false", "0", "pending"}:
                status_markers += 1
    if status_markers >= 2:
        return "keep"
    if status_markers == 1:
        return "keep"
    return ""


def _record_quality(
    *,
    scene_title: Optional[str],
    description: Optional[str],
    performers: List[str],
    categories: List[str],
    tags: List[str],
    links: List[str],
    decision: str,
    score: Optional[float],
) -> float:
    q = 0.0
    if scene_title:
        q += 0.25
    if description:
        q += 0.20
    if performers:
        q += 0.15
    if categories:
        q += 0.15
    if tags:
        q += 0.10
    if links:
        q += 0.10
    if decision:
        q += 0.03
    if score is not None:
        q += 0.02
    return round(min(q, 1.0), 3)


def _is_low_signal(row: Dict[str, Any]) -> bool:
    if row.get("decision"):
        return False
    if row.get("operator_score_100") is not None or row.get("model_score_100") is not None:
        return False
    if row.get("record_quality", 0.0) >= 0.30:
        return False
    if row.get("title") and row.get("description"):
        return False
    if row.get("filename") and row.get("source_links"):
        return False
    return True


def _fallback_scene_id(row: Dict[str, Any]) -> Optional[str]:
    sheet = row.get("_source_sheet")
    source_row = row.get("_source_row")
    if not sheet or source_row is None:
        return None
    return f"{_safe_name(str(sheet))}_r{source_row}"


def _dedupe_key(row: Dict[str, Any]) -> str:
    scene = row.get("scene_id") or ""
    filename = row.get("filename") or ""
    title = row.get("title") or ""
    source = row.get("source_file") or ""
    return f"{scene}|{filename}|{title}|{source}".lower()


def _canonical_key(key: str) -> str:
    k = str(key or "").strip().lower()
    k = k.replace("&", " and ")
    k = re.sub(r"[^a-z0-9]+", "_", k)
    k = re.sub(r"_+", "_", k)
    return k.strip("_")


def _detect_header_row(ws) -> int:
    best_row = 1
    best_score = -1
    max_scan = min(ws.max_row, 12)
    for row_idx in range(1, max_scan + 1):
        values = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
        non_empty = [str(v).strip() for v in values if v not in (None, "")]
        score = len(non_empty)
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def _row_headers(ws, row_idx: int) -> List[str]:
    headers: List[str] = []
    used: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(row=row_idx, column=c).value
        label = _canonical_key(str(raw)) if raw not in (None, "") else f"column_{c}"
        if not label:
            label = f"column_{c}"
        n = used.get(label, 0)
        used[label] = n + 1
        if n > 0:
            label = f"{label}_{n+1}"
        headers.append(label)
    return headers


def _is_empty_row(values: List[Any]) -> bool:
    for v in values:
        if v not in (None, "") and str(v).strip():
            return False
    return True


def _normalize_score(raw: Optional[str]) -> Optional[float]:
    if raw in (None, ""):
        return None
    try:
        v = float(raw)
    except (TypeError, ValueError):
        return None
    if 0 <= v <= 10:
        v *= 10.0
    if v < 0 or v > 100:
        return None
    return round(v, 1)


def _safe_name(s: str) -> str:
    out = "".join(c if c.isalnum() or c in "-_" else "_" for c in (s or "dataset"))
    return out[:80] or "dataset"
